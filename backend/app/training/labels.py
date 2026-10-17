"""Strict imports for the organizer's actual workbook and normalized companions.

Raw text timestamps support full Russian minute/second quantities only. Numeric
Excel/JSON cells are explicit seconds. MM:SS, spreadsheet dates and bare numeric
text are deliberately unsupported rather than guessed. Canonical timestamp_sec
CSV cells are explicitly numeric seconds and have a separate parser.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
import hashlib
import io
import json
import math
from pathlib import Path
import re
import unicodedata
import zipfile

SCHEMA_VERSION = "forte-labels-v1"
WORKBOOK_HEADERS = ("наименование файла", "таймстемп", "нарушение / без нарушения")
RECORD_FIELDS = ("source_sheet", "source_row", "filename_raw", "matching_key", "resolved_video_path",
                 "timestamp_raw", "timestamp_sec", "label_raw", "has_violation", "event_type",
                 "event_start_sec", "event_end_sec", "rule_id", "review_status")
VIDEO_EXT = frozenset({".mov", ".mp4", ".m4v", ".mkv", ".avi"})
MAX_INPUT_BYTES = 32 * 1024 * 1024
MAX_RECORDS = 100_000
NUMBER = r"[0-9]+(?:[.,][0-9]+)?"
MINUTE = r"(?:минута|минуты|минут|мин\.?)"
SECOND = r"(?:секунда|секунды|секунд|сек\.?)"
RUSSIAN_TIME = re.compile(
    rf"(?:(?P<minutes>{NUMBER})\s*{MINUTE}(?:\s*(?P<seconds>{NUMBER})\s*{SECOND})?"
    rf"|(?P<only_seconds>{NUMBER})\s*{SECOND})", re.IGNORECASE)


class LabelDataError(ValueError):
    """Fixed, safe metadata error; raw records belong only in local reports."""


def issue(code: str, message: str, index: int | None = None, record: dict | None = None) -> dict:
    return {"code": code, "severity": "error", "message": message, "record_index": index,
            "source_row": record.get("source_row") if record else None,
            "filename_raw": record.get("filename_raw") if record else None}


@dataclass
class ImportResult:
    records: list[dict]
    issues: list[dict] = field(default_factory=list)
    source_fingerprint: dict = field(default_factory=dict)
    input_fingerprint: dict = field(default_factory=dict)
    provenance_status: str = "input_only"
    schema_version: str = SCHEMA_VERSION

    @property
    def valid(self) -> bool:
        return not any(item["severity"] == "error" for item in self.issues)

    def to_dict(self) -> dict:
        return {"schema_version": self.schema_version, "valid": self.valid, "records": self.records,
                "issues": self.issues, "source_fingerprint": self.source_fingerprint,
                "input_fingerprint": self.input_fingerprint, "provenance_status": self.provenance_status}


def normalize_label(value) -> bool:
    if not isinstance(value, str):
        raise LabelDataError("A nonempty explicit Russian binary label is required.")
    normalized = " ".join(unicodedata.normalize("NFC", value).split()).casefold()
    mapping = {"нарушение": True, "без нарушения": False}
    if normalized not in mapping:
        raise LabelDataError("The label is missing or is not an exact supported binary label.")
    return mapping[normalized]


def _nonnegative_number(value) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise LabelDataError("The timestamp must be a finite nonnegative number.")
    number = float(value)
    if not math.isfinite(number) or number < 0:
        raise LabelDataError("The timestamp must be a finite nonnegative number.")
    return number


def parse_timestamp(value) -> float | None:
    """Parse complete Russian text or a numeric seconds cell; retain null blanks."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _nonnegative_number(value)
    if not isinstance(value, str):
        raise LabelDataError("Unsupported timestamp type; spreadsheet times are not assumed to be seconds.")
    match = RUSSIAN_TIME.fullmatch(unicodedata.normalize("NFC", value).strip())
    if not match:
        raise LabelDataError("Unsupported timestamp text; use complete Russian minute/second units.")
    minute = float((match.group("minutes") or "0").replace(",", "."))
    second = float((match.group("seconds") or match.group("only_seconds") or "0").replace(",", "."))
    return _nonnegative_number(minute * 60 + second)


def parse_numeric_seconds(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        if not re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", value.strip()):
            raise LabelDataError("Canonical timestamp_sec must contain numeric seconds or be empty.")
        value = float(value)
    return _nonnegative_number(value)


def matching_key(name: str) -> str:
    """Known leading Х/X alias only; never transliterate the rest of a filename."""
    name = str(name)
    suffix = Path(name).suffix
    stem = name[:-len(suffix)] if suffix.lower() in VIDEO_EXT else name
    return "X" + stem[1:] if stem.startswith("Х") else stem


def fingerprint(path: Path, data: bytes | None = None) -> dict:
    content = data if data is not None else path.read_bytes()
    return {"sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content),
            "source_name": path.name, "format": path.suffix.lower().lstrip(".")}


def _nullable(value):
    return None if value is None or value == "" else value


def _explicit_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().casefold() in {"true", "false"}:
        return value.strip().casefold() == "true"
    raise LabelDataError("Normalized has_violation must explicitly be true or false.")


def normalize_record(row: dict, index: int, *, native: bool = False) -> tuple[dict, list[dict]]:
    errors = []
    if native:
        row = {"filename_raw": row.get(WORKBOOK_HEADERS[0]), "timestamp_raw": row.get(WORKBOOK_HEADERS[1]),
               "label_raw": row.get(WORKBOOK_HEADERS[2]), "source_sheet": row.get("source_sheet", "labels"),
               "source_row": row.get("source_row", index + 2)}
    filename = row.get("filename_raw")
    record = {"source_sheet": _nullable(row.get("source_sheet")),
              "source_row": row.get("source_row", index + 2),
              "filename_raw": filename, "matching_key": None, "resolved_video_path": None,
              "timestamp_raw": row.get("timestamp_raw"), "timestamp_sec": None,
              "label_raw": row.get("label_raw"), "has_violation": None,
              "event_type": _nullable(row.get("event_type")),
              "event_start_sec": _nullable(row.get("event_start_sec")),
              "event_end_sec": _nullable(row.get("event_end_sec")),
              "rule_id": _nullable(row.get("rule_id")),
              "review_status": _nullable(row.get("review_status")) or "not_reviewed"}
    try:
        source_row = record["source_row"]
        if isinstance(source_row, str) and source_row.isascii() and source_row.isdigit():
            source_row = int(source_row)
        if isinstance(source_row, bool) or not isinstance(source_row, int) or source_row < 1:
            raise ValueError
        record["source_row"] = source_row
    except (ValueError, TypeError):
        errors.append(issue("invalid_source_row", "Source row must be a positive integer.", index, record))
    if (not isinstance(filename, str) or not filename.strip() or any(char in filename for char in "/\\\x00")
            or filename in {".", ".."}):
        errors.append(issue("invalid_identifier", "A nonempty filename identifier without path components is required.", index, record))
    else:
        record["matching_key"] = matching_key(filename)
        supplied_key = _nullable(row.get("matching_key"))
        if supplied_key is not None and supplied_key != record["matching_key"]:
            errors.append(issue("invalid_identifier", "The supplied matching key uses an unsupported filename transformation.", index, record))
    try:
        record["has_violation"] = normalize_label(record["label_raw"])
    except LabelDataError as exc:
        errors.append(issue("missing_or_unknown_label", str(exc), index, record))
    if "has_violation" in row:
        try:
            declared = _explicit_bool(row["has_violation"])
            if record["has_violation"] is not None and declared != record["has_violation"]:
                raise LabelDataError("Normalized has_violation disagrees with the original Russian label.")
        except LabelDataError as exc:
            record["has_violation"] = None
            errors.append(issue("missing_or_unknown_label", str(exc), index, record))
    try:
        record["timestamp_sec"] = parse_timestamp(record["timestamp_raw"])
        if "timestamp_sec" in row:
            declared = parse_numeric_seconds(row["timestamp_sec"])
            if declared != record["timestamp_sec"]:
                raise LabelDataError("Normalized timestamp_sec disagrees with timestamp_raw.")
        if record["has_violation"] is True and record["timestamp_sec"] is None:
            raise LabelDataError("An explicitly positive row requires a timestamp.")
        if record["has_violation"] is False and record["timestamp_sec"] is not None:
            raise LabelDataError("An explicitly negative row must have a blank timestamp; zero is not blank.")
    except LabelDataError as exc:
        errors.append(issue("invalid_timestamp", str(exc), index, record))
    # Companions are source metadata. They cannot assert a successful disk resolution.
    # A fresh local resolver always establishes the actual path and method.
    if _nullable(row.get("resolved_video_path")) is not None:
        record["previous_resolved_video_path"] = row["resolved_video_path"]
    if (any(record[field] is not None for field in ("event_type", "event_start_sec", "event_end_sec", "rule_id"))
            or record["review_status"] != "not_reviewed"):
        errors.append(issue("unsupported_source_annotation", "Organizer source metadata has no reviewed events, boundaries or rule mappings; store annotations separately.", index, record))
    return record, errors


def _read_workbook(path: Path, content: bytes) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise LabelDataError("XLSX imports require the locally installed openpyxl dependency.") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 64 * 1024 * 1024:
                raise LabelDataError("The workbook exceeds the supported uncompressed size.")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if "labels" not in workbook.sheetnames:
                raise LabelDataError("The workbook must contain the exact worksheet name labels.")
            rows = workbook["labels"].iter_rows(values_only=True)
            header = next(rows, ())
            if tuple(header[:3]) != WORKBOOK_HEADERS or any(cell is not None for cell in header[3:]):
                raise LabelDataError("The labels worksheet must have the three exact organizer headers in A1:C1.")
            output = []
            for row_number, values in enumerate(rows, 2):
                if row_number > MAX_RECORDS + 1:
                    raise LabelDataError("The label worksheet exceeds the supported row range.")
                if not any(value is not None for value in values):
                    continue
                if len(output) >= MAX_RECORDS:
                    raise LabelDataError("The label input exceeds the supported row count.")
                if any(value is not None for value in values[3:]):
                    raise LabelDataError("Unexpected populated columns occur outside the three organizer columns.")
                row = dict(zip(WORKBOOK_HEADERS, values[:3]))
                row.update(source_sheet="labels", source_row=row_number)
                output.append(row)
            return output
        finally:
            workbook.close()
    except LabelDataError:
        raise
    except Exception as exc:
        raise LabelDataError("The XLSX workbook could not be read safely.") from exc


def import_labels(path: str | Path, *, source_workbook: str | Path | None = None) -> ImportResult:
    path = Path(path)
    result = ImportResult(records=[])
    try:
        if path.stat().st_size > MAX_INPUT_BYTES:
            raise LabelDataError("The label file exceeds the supported size limit.")
        content = path.read_bytes()
        result.input_fingerprint = fingerprint(path, content)
        result.source_fingerprint = dict(result.input_fingerprint)
        native = False
        if path.suffix.lower() == ".xlsx":
            rows = _read_workbook(path, content)
            native = True
            result.provenance_status = "original_workbook"
        elif path.suffix.lower() == ".json":
            data = json.loads(content.decode("utf-8-sig"))
            if not isinstance(data, dict) or not isinstance(data.get("records"), list):
                raise LabelDataError("Normalized JSON must be an object with a top-level records array.")
            rows = data["records"]
            claimed = data.get("source_fingerprint")
            if claimed is not None:
                if not isinstance(claimed, dict) or not re.fullmatch(r"[a-fA-F0-9]{64}", str(claimed.get("sha256", ""))):
                    raise LabelDataError("The supplied source fingerprint is malformed.")
                result.source_fingerprint = dict(claimed)
                result.provenance_status = "claimed_source_unverified"
        elif path.suffix.lower() == ".csv":
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig"), newline=""))
            headers = reader.fieldnames or []
            required = {"filename_raw", "timestamp_raw", "timestamp_sec", "label_raw", "has_violation"}
            if len(set(headers)) != len(headers) or not required.issubset(headers):
                raise LabelDataError("Normalized CSV requires exact canonical column names including original raw fields.")
            rows = list(reader)
            if any(None in row for row in rows):
                raise LabelDataError("CSV rows contain more fields than the declared header.")
        else:
            raise LabelDataError("Supported label formats are .xlsx, normalized .json and UTF-8 BOM .csv.")
        if len(rows) > MAX_RECORDS:
            raise LabelDataError("The label input exceeds the supported row count.")
        for index, row in enumerate(rows):
            if not isinstance(row, dict):
                result.issues.append(issue("invalid_format", "Each normalized record must be an object.", index))
                result.records.append({"source_row": index + 2, "filename_raw": None, "has_violation": None,
                                       "timestamp_raw": None, "timestamp_sec": None, "resolved_video_path": None})
                continue
            record, errors = normalize_record(row, index, native=native)
            result.records.append(record)
            result.issues.extend(errors)
        if not rows:
            result.issues.append(issue("invalid_format", "The label input contains no records."))
        seen = {}
        for index, record in enumerate(result.records):
            name = record.get("filename_raw")
            if isinstance(name, str) and name in seen:
                result.issues.append(issue("duplicate_identifier", "A raw video identifier occurs in multiple label records; explicit review is required.", index, record))
                result.issues.append(issue("duplicate_identifier", "A raw video identifier occurs in multiple label records; explicit review is required.", seen[name], result.records[seen[name]]))
            elif isinstance(name, str):
                seen[name] = index
        if source_workbook is not None:
            reference = import_labels(source_workbook)
            if Path(source_workbook).suffix.lower() != ".xlsx" or not reference.valid:
                result.issues.append(issue("source_mismatch", "The source workbook is unavailable or not a valid organizer worksheet."))
            else:
                claimed = result.source_fingerprint if result.provenance_status == "claimed_source_unverified" else None
                if claimed and claimed["sha256"].lower() != reference.source_fingerprint["sha256"]:
                    result.issues.append(issue("source_mismatch", "The supplied source fingerprint does not match the workbook."))
                def signature(records):
                    values = [(row.get("filename_raw"), row.get("source_sheet"), row.get("source_row"),
                               _nullable(row.get("timestamp_raw")), row.get("timestamp_sec"),
                               row.get("label_raw"), row.get("has_violation")) for row in records]
                    return sorted(values, key=lambda value: (str(value[1]), str(value[2]), str(value[0])))
                if signature(result.records) != signature(reference.records):
                    result.issues.append(issue("source_mismatch", "Normalized records do not exactly match the original workbook raw fields and labels."))
                result.source_fingerprint = reference.source_fingerprint
                result.provenance_status = "verified_against_source" if not any(
                    item["code"] == "source_mismatch" for item in result.issues) else "source_mismatch"
    except (OSError, UnicodeError, ValueError, csv.Error) as exc:
        message = str(exc) if isinstance(exc, LabelDataError) else "The label input could not be read as the declared format."
        result.issues.append(issue("invalid_format", message))
    return result
